"""API tests: file upload / preview / listing + artifact session attribution."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from ginno_runtime.files import reset_registries
from ginno_runtime.testing.fake_model import script

pytestmark = pytest.mark.api


@pytest.fixture(autouse=True)
def _fresh_files(isolated_home):
    reset_registries()
    yield
    reset_registries()


@pytest.fixture
def sid(client, patch_build_model, tmp_path):
    patch_build_model(script(text="ok"))
    ws = tmp_path / "ws"
    ws.mkdir()
    r = client.post(
        "/api/sessions",
        json={"project_slug": "default", "workspace": str(ws), "agent_id": "dev"},
    ).json()
    assert r["ok"] is True
    return r["id"]


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["产品", "金额"])
    ws.append(["A", 10.0])
    ws.append(["B", 20.0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("合同正文内容")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_xlsx_registers_file_and_artifact(client, sid, tmp_path):
    r = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("报表.xlsx", _xlsx_bytes(), "application/vnd.ms-excel")},
    ).json()
    assert r["ok"] is True
    f = r["file"]
    assert f["kind"] == "spreadsheet"
    assert f["session_id"] == sid
    assert f["artifact_id"]
    # landed under the session files dir: projects/<slug>/sessions/<sid>/uploads/
    assert f"/sessions/{sid}/uploads/" in f["path"]
    assert Path(f["path"]).is_file()
    # artifact created WITH session attribution (the §7.6 fix)
    arts = client.get("/api/artifacts?project_slug=default").json()
    art = next(a for a in arts if a["id"] == f["artifact_id"])
    assert art["kind"] == "file"
    assert art["session_id"] == sid
    assert art["ref"] == f["path"]


def test_upload_unknown_session_404(client):
    r = client.post(
        "/api/files",
        data={"session_id": "nope"},
        files={"file": ("a.csv", b"a,b\n1,2\n", "text/csv")},
    )
    assert r.status_code == 404


def test_upload_sanitizes_filename(client, sid):
    r = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("../../evil/..\\x.xlsx", _xlsx_bytes(), "application/octet-stream")},
    ).json()
    assert r["ok"] is True
    name = r["file"]["name"]
    assert "/" not in name and "\\" not in name and ".." not in name
    assert Path(r["file"]["path"]).name.endswith(name)


def test_preview_spreadsheet(client, sid):
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("报表.xlsx", _xlsx_bytes(), "application/octet-stream")},
    ).json()
    pv = client.get(f"/api/files/{up['file']['id']}/preview").json()
    assert pv["ok"] is True
    assert pv["kind"] == "spreadsheet"
    assert pv["sheets"] == [{"name": "销售", "rows": 2, "cols": 2}]
    assert pv["columns"][0]["name"] == "产品"
    assert pv["rows"][0][0] == "A"
    assert float(pv["rows"][0][1]) == 10.0  # calamine may read integral cells as int
    assert pv["total_rows"] == 2
    assert pv["file"]["name"] == "报表.xlsx"


def test_preview_pagination_params(client, sid):
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("t.csv", b"a\n1\n2\n3\n", "text/csv")},
    ).json()
    pv = client.get(f"/api/files/{up['file']['id']}/preview?offset=1&limit=1").json()
    assert pv["ok"] is True
    assert pv["rows"] == [["2"]]
    assert pv["offset"] == 1 and pv["total_rows"] == 3


def test_preview_docx_returns_markdown(client, sid):
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("c.docx", _docx_bytes(), "application/octet-stream")},
    ).json()
    pv = client.get(f"/api/files/{up['file']['id']}/preview").json()
    assert pv["ok"] is True
    assert pv["kind"] == "document"
    assert "合同正文内容" in pv["markdown"]


def test_preview_unsupported_format(client, sid):
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("x.exe", b"MZ\x90\x00", "application/octet-stream")},
    ).json()
    assert up["ok"] is True  # upload itself is format-agnostic
    pv = client.get(f"/api/files/{up['file']['id']}/preview").json()
    assert pv["ok"] is False
    assert "不支持" in pv["error"]


def test_preview_unknown_id(client):
    pv = client.get("/api/files/deadbeef00/preview").json()
    assert pv["ok"] is False and "not found" in pv["error"]


def test_list_files_scoped_by_session(client, sid, patch_build_model, tmp_path):
    patch_build_model(script(text="ok"))
    ws2 = tmp_path / "ws2"
    ws2.mkdir()
    sid2 = client.post(
        "/api/sessions",
        json={"project_slug": "default", "workspace": str(ws2), "agent_id": "dev"},
    ).json()["id"]

    client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("a.csv", b"a\n1\n", "text/csv")},
    )
    client.post(
        "/api/files",
        data={"session_id": sid2},
        files={"file": ("b.csv", b"b\n2\n", "text/csv")},
    )
    mine = client.get(f"/api/files?project_slug=default&session_id={sid}").json()
    assert [f["name"] for f in mine] == ["a.csv"]
    all_ = client.get("/api/files?project_slug=default").json()
    assert {f["name"] for f in all_} == {"a.csv", "b.csv"}


# --------------------------------------------------------------------------
# download / export
# --------------------------------------------------------------------------

def test_download_raw_returns_original_bytes(client, sid):
    body = "名称,金额\n甲,1\n".encode()
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("报表.csv", body, "text/csv")},
    ).json()
    r = client.get(f"/api/files/{up['file']['id']}/download")
    assert r.status_code == 200
    assert r.content == body
    cd = r.headers["content-disposition"]
    assert "attachment" in cd
    assert "filename*=UTF-8''" in cd and "%E6%8A%A5%E8%A1%A8" in cd  # 报表 percent-encoded


def test_download_csv_export_from_xlsx(client, sid):
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("报表.xlsx", _xlsx_bytes(), "application/octet-stream")},
    ).json()
    r = client.get(f"/api/files/{up['file']['id']}/download?fmt=csv&sheet=销售")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    text = r.content.decode("utf-8-sig")
    assert "产品,金额" in text and "A,10" in text


def test_download_csv_export_rejects_docx(client, sid):
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("c.docx", _docx_bytes(), "application/octet-stream")},
    ).json()
    r = client.get(f"/api/files/{up['file']['id']}/download?fmt=csv")
    assert r.status_code == 400
    assert "表格" in r.json()["detail"]


def test_download_unknown_id_404(client):
    r = client.get("/api/files/deadbeef00/download")
    assert r.status_code == 404


def test_save_to_downloads_raw_and_dedupe(client, sid, monkeypatch, tmp_path):
    dl = tmp_path / "downloads"
    monkeypatch.setenv("GINNO_DOWNLOADS", str(dl))
    body = b"a,b\n1,2\n"
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("t.csv", body, "text/csv")},
    ).json()
    fid = up["file"]["id"]

    r1 = client.post(f"/api/files/{fid}/save-to-downloads", json={"fmt": "raw"}).json()
    assert r1["ok"] is True and r1["name"] == "t.csv"
    assert (dl / "t.csv").read_bytes() == body

    # second save must not clobber → " (1)" suffix
    r2 = client.post(f"/api/files/{fid}/save-to-downloads", json={}).json()
    assert r2["ok"] is True and r2["name"] == "t (1).csv"
    assert (dl / "t (1).csv").is_file()


def test_save_to_downloads_csv_export(client, sid, monkeypatch, tmp_path):
    dl = tmp_path / "downloads"
    monkeypatch.setenv("GINNO_DOWNLOADS", str(dl))
    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("报表.xlsx", _xlsx_bytes(), "application/octet-stream")},
    ).json()
    r = client.post(
        f"/api/files/{up['file']['id']}/save-to-downloads",
        json={"fmt": "csv", "sheet": "销售"},
    ).json()
    assert r["ok"] is True and r["name"] == "报表.csv"  # single sheet → no suffix
    text = (dl / "报表.csv").read_bytes().decode("utf-8-sig")
    assert "产品,金额" in text


def test_save_to_downloads_unknown_id(client):
    r = client.post("/api/files/deadbeef00/save-to-downloads", json={}).json()
    assert r["ok"] is False and "not found" in r["error"]


# ------------------- artifact metadata inspector ------------------- #
def test_artifact_metadata_and_user_corrections(client, sid):
    from ginno_runtime import server

    up = client.post(
        "/api/files",
        data={"session_id": sid},
        files={"file": ("报表.xlsx", _xlsx_bytes(), "application/octet-stream")},
    ).json()
    file_id = up["file"]["id"]
    arts = client.get("/api/artifacts?project_slug=default").json()
    art = next(a for a in arts if a["kind"] == "file")

    # Inspector payload: exact injectable schema + provenance + file facts
    meta = client.get(f"/api/artifacts/{art['id']}/metadata").json()
    assert meta["ok"] is True and meta["exists"] is True
    assert meta["schema_source"] == "computed"
    assert "产品" in meta["schema"] and meta["file"]["id"] == file_id

    # User corrections round-trip: rename + schema override + kind fix
    r = client.put(
        f"/api/artifacts/{art['id']}",
        json={"name": "2026销售报表", "schema": "[销售] 2行×2列（人工修正）", "file_kind": "table"},
    ).json()
    assert r["ok"] is True and r["artifact"]["name"] == "2026销售报表"
    meta = client.get(f"/api/artifacts/{art['id']}/metadata").json()
    assert meta["schema"] == "[销售] 2行×2列（人工修正）"
    assert meta["schema_source"] == "override"
    assert meta["file"]["kind"] == "table"  # registry correction persisted

    # The override (not the recomputed summary) is what injection now uses
    item = server._resolve_attached_files([{"id": file_id}], "default", sid)[0]
    assert item["schema"] == "[销售] 2行×2列（人工修正）"

    # Foolproof guards: unknown ids, blank names
    assert client.get("/api/artifacts/nope/metadata").json()["ok"] is False
    assert client.put("/api/artifacts/nope", json={"name": "x"}).json()["ok"] is False
    bad = client.put(f"/api/artifacts/{art['id']}", json={"name": "   "}).json()
    assert bad["ok"] is False
    assert client.get(f"/api/artifacts/{art['id']}/metadata").json()["artifact"]["name"] == "2026销售报表"
