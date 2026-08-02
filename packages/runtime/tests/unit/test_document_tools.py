"""Unit tests: parse_document / analyze_table tools."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from ginno_runtime.tools.document_tools import analyze_table, parse_document

pytestmark = pytest.mark.unit


def make_csv(p: Path) -> Path:
    p.write_text("a,b\n1,2\n3,4\n5,6\n", encoding="utf-8")
    return p


def make_xlsx(p: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["地区", "金额"])
    ws.append(["北", 100])
    ws.append(["南", 300])
    ws.append(["北", 50])
    wb.save(p)
    return p


def make_docx(p: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.core_properties.author = "张三"
    doc.add_paragraph("协议内容")
    doc.save(p)
    return p


# --------------------------------------------------------------------------- #
# parse_document
# --------------------------------------------------------------------------- #

def test_parse_document_text(tmp_path):
    f = make_docx(tmp_path / "c.docx")
    out = parse_document.invoke({"path": str(f)})
    assert "协议内容" in out


def test_parse_document_metadata(tmp_path):
    f = make_docx(tmp_path / "c.docx")
    d = json.loads(parse_document.invoke({"path": str(f), "format": "metadata"}))
    assert d["ok"] is True
    assert d["kind"] == "document"
    assert d["metadata"]["author"] == "张三"


def test_parse_document_json(tmp_path):
    f = make_csv(tmp_path / "t.csv")
    d = json.loads(parse_document.invoke({"path": str(f), "format": "json"}))
    assert d["ok"] is True and d["kind"] == "table"
    assert "| a | b |" in d["text"]


def test_parse_document_errors_are_json(tmp_path):
    d = json.loads(parse_document.invoke({"path": str(tmp_path / "nope.xlsx")}))
    assert d["ok"] is False
    f = tmp_path / "x.exe"
    f.write_bytes(b"MZ")
    d2 = json.loads(parse_document.invoke({"path": str(f)}))
    assert d2["ok"] is False and "不支持" in d2["error"]


# --------------------------------------------------------------------------- #
# analyze_table
# --------------------------------------------------------------------------- #

def test_analyze_scalar_result(tmp_path):
    f = make_csv(tmp_path / "t.csv")
    d = json.loads(analyze_table.invoke({"path": str(f), "code": "result = df['b'].sum()"}))
    assert d["ok"] is True
    assert d["result_type"] != "table"
    assert int(d["result"]) == 12


def test_analyze_table_result_with_derived_csv(tmp_path):
    f = make_xlsx(tmp_path / "s.xlsx")
    code = (
        "result = df.groupby('地区')['金额'].sum().reset_index()"
        ".sort_values('金额', ascending=False)"
    )
    d = json.loads(analyze_table.invoke({"path": str(f), "code": code}))
    assert d["ok"] is True
    assert d["result_type"] == "table"
    cols = d["result"]["columns"]
    assert "地区" in cols and "金额" in cols
    first = d["result"]["rows"][0]
    assert first[cols.index("地区")] == "南"  # 300 > 150
    # derived CSV written next to the source for auto-preview
    assert "derived_path" in d
    derived = Path(d["derived_path"])
    assert derived.is_file() and derived.suffix == ".csv"
    assert "results" in derived.parts


def test_analyze_syntax_error_reports(tmp_path):
    f = make_csv(tmp_path / "t.csv")
    d = json.loads(analyze_table.invoke({"path": str(f), "code": "result = df["}))
    assert d["ok"] is False
    assert "stderr" in d


def test_analyze_timeout(tmp_path):
    f = make_csv(tmp_path / "t.csv")
    d = json.loads(
        analyze_table.invoke(
            {"path": str(f), "code": "import time\ntime.sleep(9)\nresult = 1", "timeout": 2}
        )
    )
    assert d["ok"] is False and "超时" in d["error"]


def test_analyze_rejects_non_table(tmp_path):
    f = make_docx(tmp_path / "c.docx")
    d = json.loads(analyze_table.invoke({"path": str(f), "code": "result = 1"}))
    assert d["ok"] is False and "仅支持" in d["error"]


def test_analyze_missing_file(tmp_path):
    d = json.loads(
        analyze_table.invoke({"path": str(tmp_path / "gone.csv"), "code": "result = 1"})
    )
    assert d["ok"] is False and "不存在" in d["error"]


def test_analyze_load_error_inside_sandbox(tmp_path):
    f = tmp_path / "broken.csv"
    f.write_bytes(b"\xff\xfe\x00bad")
    d = json.loads(analyze_table.invoke({"path": str(f), "code": "result = 1"}))
    # pandas may still parse garbage as one column; accept either outcome but
    # it must be well-formed JSON with ok flag
    assert isinstance(d.get("ok"), bool)
